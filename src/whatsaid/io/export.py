"""Excel export for extracted resources."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from ..core.db import DEFAULT_DB_PATH, get_connection

_HEADERS = ["Chat", "Date", "Person", "Tags", "Title", "Platform", "URL", "Context", "Status"]

_COL_URL = "G"
_COL_CONTEXT = "H"
_URL_WIDTH = 50
_CONTEXT_WIDTH = 60


def generate_excel(
    output_path: str = "Extracted_Resources.xlsx",
    db_path: str = DEFAULT_DB_PATH,
    chat_id: int | None = None,
) -> int:
    """Fetch resources from the DB and write them to an Excel workbook.

    Args:
        output_path: Destination ``.xlsx`` file path.
        db_path:     SQLite database path.
        chat_id:     If given, export only resources from this chat.

    Returns the number of rows exported.
    """
    conn = get_connection(db_path)
    cursor = conn.cursor()

    if chat_id is not None:
        cursor.execute("""
            SELECT
                ch.name          AS chat,
                m.timestamp      AS date,
                m.sender,
                r.tags,
                r.title,
                r.platform,
                r.canonical_url  AS url,
                r.context,
                r.status
            FROM resources r
            JOIN messages m  ON r.first_message_id = m.id
            JOIN chats    ch ON r.chat_id = ch.id
            WHERE r.chat_id = ?
        """, (chat_id,))
    else:
        cursor.execute("""
            SELECT
                ch.name          AS chat,
                m.timestamp      AS date,
                m.sender,
                r.tags,
                r.title,
                r.platform,
                r.canonical_url  AS url,
                r.context,
                r.status
            FROM resources r
            JOIN messages m  ON r.first_message_id = m.id
            JOIN chats    ch ON r.chat_id = ch.id
            ORDER BY ch.name, m.timestamp
        """)

    rows = cursor.fetchall()
    conn.close()

    wb = Workbook()

    # ---- Resources sheet ------------------------------------------------
    ws_resources = wb.active
    ws_resources.title = "Resources"

    ws_resources.append(_HEADERS)
    for cell in ws_resources[1]:
        cell.font = Font(bold=True)

    platform_counts: dict[str, int] = {}
    chat_counts: dict[str, int] = {}

    for row in rows:
        platform = row["platform"]
        chat_name = row["chat"]
        platform_counts[platform] = platform_counts.get(platform, 0) + 1
        chat_counts[chat_name] = chat_counts.get(chat_name, 0) + 1

        ws_resources.append([
            chat_name,
            row["date"],
            row["sender"],
            row["tags"] or "",
            row["title"] or "",
            platform,
            row["url"],
            row["context"],
            row["status"],
        ])

    ws_resources.column_dimensions[_COL_URL].width = _URL_WIDTH
    ws_resources.column_dimensions[_COL_CONTEXT].width = _CONTEXT_WIDTH
    for cell in ws_resources[_COL_CONTEXT]:
        cell.alignment = Alignment(wrap_text=True)

    # ---- Summary sheet --------------------------------------------------
    ws_summary = wb.create_sheet("Summary")
    ws_summary.append(["Extracted Resources Summary"])
    ws_summary["A1"].font = Font(bold=True, size=14)
    ws_summary.append([])
    ws_summary.append(["Total unique resources:", len(rows)])
    ws_summary.append([])

    ws_summary.append(["By Chat"])
    ws_summary.cell(row=ws_summary.max_row, column=1).font = Font(bold=True)
    for chat_name, count in sorted(chat_counts.items()):
        ws_summary.append([chat_name, count])

    ws_summary.append([])
    ws_summary.append(["Platform Breakdown"])
    ws_summary.cell(row=ws_summary.max_row, column=1).font = Font(bold=True)
    for platform, count in sorted(platform_counts.items(), key=lambda x: x[1], reverse=True):
        ws_summary.append([platform, count])

    wb.save(output_path)
    print(f"Exported {len(rows)} resources to {output_path}")
    return len(rows)
